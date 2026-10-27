"""End-to-end test: a file that failed verification must NOT be counted as downloaded.

Scenario:
  videos/good-slug/Good [111111].mp4   -> valid, has audio
  videos/bad-slug/Bad  [222222].mp4   -> corrupt (no audio track)
  report.json says bad-slug failed

Expected after build_index_xlsx.py:
  good-slug -> 是否已下载 = 是
  bad-slug  -> 是否已下载 = 否
  KPI "已下载到本地" == 1 (not 2)
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

SKILL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(SKILL, "scripts")
ROOT = tempfile.mkdtemp(prefix="e2e_")
VIDEOS = os.path.join(ROOT, "videos")
os.makedirs(os.path.join(VIDEOS, "good-slug"))
os.makedirs(os.path.join(VIDEOS, "bad-slug"))

# locate ffmpeg
ffmpeg = shutil.which("ffmpeg")
if not ffmpeg:
    print("SKIP: ffmpeg not found")
    sys.exit(0)

subprocess.run([ffmpeg, "-y", "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=10",
                "-f", "lavfi", "-i", "sine=frequency=440:duration=2",
                "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
                os.path.join(VIDEOS, "good-slug", "Good [111111].mp4")],
               capture_output=True)
subprocess.run([ffmpeg, "-y", "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=10",
                "-c:v", "libx264", "-pix_fmt", "yuv420p",
                os.path.join(VIDEOS, "bad-slug", "Bad [222222].mp4")],
               capture_output=True)

# manifest with two events whose vimeo ids match the filenames
manifest = {
    "listing": "test", "base": "test", "count": 2, "with_video": 2,
    "skipped_hub_pages": 0, "errors": [],
    "events": [
        {"slug": "good-slug", "title": "Good Video", "url": "https://x/1", "date": "2026-01-01",
         "time": "", "videos": [{"platform": "vimeo", "id": "111111"}], "has_video": True,
         "presenters": [], "attachments": [], "description": "good one"},
        {"slug": "bad-slug", "title": "Bad Video", "url": "https://x/2", "date": "2026-01-02",
         "time": "", "videos": [{"platform": "vimeo", "id": "222222"}], "has_video": True,
         "presenters": [], "attachments": [], "description": "bad one"},
    ],
}
mpath = os.path.join(ROOT, "manifest.json")
with open(mpath, "w", encoding="utf-8") as f:
    json.dump(manifest, f, ensure_ascii=False)

# verification report: bad-slug failed
rpath = os.path.join(ROOT, "report.json")
with open(rpath, "w", encoding="utf-8") as f:
    json.dump({"ok": 1, "skipped": 0, "failed": 1,
               "failures": [{"slug": "bad-slug", "error": "verify failed"}],
               "videos": {"good-slug": {"ok": True, "file": "x", "report": {}},
                          "bad-slug": {"ok": False, "file": "y", "report": {"error": "no audio"}}}},
              f, ensure_ascii=False, indent=1)

xlsx = os.path.join(ROOT, "index.xlsx")
r = subprocess.run([sys.executable, os.path.join(SCRIPTS, "build_index_xlsx.py"),
                    "--manifest", mpath, "--out", xlsx, "--title", "E2E",
                    "--report", rpath, "--videos-dir", VIDEOS],
                   capture_output=True, text=True)
print(r.stdout.strip())
if r.returncode != 0:
    print("build_index FAILED:", r.stderr[-800:])
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

wb = openpyxl.load_workbook(xlsx, data_only=False)
d = wb["明细"]
res = {}
for r_ in range(2, 10):
    title = d.cell(r_, 3).value
    if title:
        res[title] = (d.cell(r_, 14).value, d.cell(r_, 13).value)
sm = wb["汇总"]

# formula cells need recalc to read values; read the formula text instead
b8_formula = sm.cell(8, 2).value

print()
print("明细 '是否已下载':")
for title, (flag, path) in res.items():
    print("   %-12s -> %s   (%s)" % (title, flag, path))
print()
print("汇总 B8 公式:", b8_formula)

ok = True
if res.get("Good Video", (None,))[0] != "是":
    print("FAIL: Good Video 应标为已下载")
    ok = False
if res.get("Bad Video", (None,))[0] != "否":
    print("FAIL: Bad Video 校验失败，不应标为已下载")
    ok = False
print()
print("RESULT:", "PASS" if ok else "FAIL")
sys.exit(0 if ok else 1)
