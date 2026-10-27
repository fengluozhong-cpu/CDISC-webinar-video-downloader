"""Build an Excel index from a harvested event manifest.

Consumes the JSON produced by harvest_events.py and produces a multi-sheet
workbook: 汇总 (rollups) / 专题 (filtered subset) / 下载说明 (provenance) /
明细 (full table).  Statistics are live Excel formulas referencing 明细, so
they stay correct if rows are edited.

Usage
-----
    python build_index_xlsx.py --manifest manifest.json --out index.xlsx \\
        --title "CDISC Public Webinars 视频清单" --keyword "SDTM" --subset "SDTM 专题"
"""
import argparse
import json
import os
import re
from datetime import date, datetime

try:
    import openpyxl
except ImportError:
    import subprocess
    import sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def xl_color(css_hex):
    v = css_hex.removeprefix("#").upper()
    if len(v) != 6:
        raise ValueError("expected #RRGGBB, got %s" % css_hex)
    return "FF" + v


XL_PRIMARY = xl_color("#4472C4")
XL_LIGHT = xl_color("#D9E2F3")
XL_DARK = xl_color("#2F5597")
XL_BORDER = xl_color("#BFBFBF")
XL_WHITE = xl_color("#FFFFFF")
XL_OK_BG = xl_color("#C6EFCE")
XL_OK_FG = xl_color("#006100")
XL_WARN_BG = xl_color("#FFC7CE")
XL_WARN_FG = xl_color("#9C0006")
XL_NOTE = xl_color("#F2F2F2")

THIN = Side(style="thin", color=XL_BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TOPIC_RULES = [
    ("SDTM", ["sdtm"]), ("SEND", ["send"]), ("ADaM", ["adam"]), ("TMF", ["tmf"]),
    ("术语 Terminology", ["terminology", "controlled terminology"]),
    ("360i", ["360i"]), ("Dataset-JSON", ["dataset-json"]),
    ("USDM / 研究设计", ["usdm", "schedule of activities"]),
    ("AI 创新", ["ai innovation", "artificial intelligence", "ai challenge"]),
]


def topic_of(row):
    t = (row.get("title") or "").lower()
    d = (row.get("description") or "").lower()
    for label, keys in TOPIC_RULES:
        if any(k in t for k in keys):
            return label
    for label, keys in TOPIC_RULES:
        if any(k in d for k in keys):
            return label
    return "其他"


def video_cell(row):
    """Human-readable video link + id."""
    vids = row.get("videos") or []
    if not vids:
        return "", ""
    v = vids[0]
    if v["platform"] == "vimeo":
        return "https://player.vimeo.com/video/%s" % v["id"], v["id"]
    if v["platform"] == "youtube":
        return "https://www.youtube.com/watch?v=%s" % v["id"], v["id"]
    if v["platform"] == "wistia":
        return "https://fast.wistia.net/embed/iframe/%s" % v["id"], v["id"]
    return v["id"], v["id"]


def parse_dt(iso, tm):
    if not iso:
        return None
    try:
        d = datetime.strptime(iso, "%Y-%m-%d").date()
    except ValueError:
        return None
    if tm:
        for fmt in ("%I:%M%p", "%I:%M %p", "%H:%M"):
            try:
                return datetime.combine(d, datetime.strptime(tm.strip(), fmt).time())
            except ValueError:
                continue
    return d


def main():
    ap = argparse.ArgumentParser(description="Build an xlsx index from an event manifest")
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--out", default="index.xlsx")
    ap.add_argument("--title", default="Webinar 视频清单")
    ap.add_argument("--keyword", default=None, help="subset filter keyword, e.g. SDTM")
    ap.add_argument("--subset-name", default="专题")
    ap.add_argument("--videos-dir", default="videos", help="folder to detect local downloads")
    ap.add_argument("--notes", default=None, help="path to a text file with provenance notes")
    ap.add_argument("--report", default=None,
                    help="verification report written by download_webinar.py --report. "
                         "Files marked failed there are NOT counted as downloaded.")
    args = ap.parse_args()

    data = json.load(open(args.manifest, encoding="utf-8"))
    events = data.get("events") or data.get("rows") or []
    if not events:
        raise SystemExit("manifest contains no events")
    kw = args.keyword or ""

    # Match local downloads by video id so the "downloaded" column stays truthful.
    # Walk RECURSIVELY: download_webinar.py --manifest writes each video into
    # videos/<slug>/, so a flat listdir() would miss every batch download and
    # silently report "not downloaded".
    #
    # A verification report (download_webinar.py --report) overrides raw file
    # presence: files listed as failed are quarantined by the downloader and
    # must NOT be counted as downloaded.
    verified = {}
    if args.report and os.path.exists(args.report):
        with open(args.report, encoding="utf-8") as f:
            rep = json.load(f)
        for slug, info in (rep.get("videos") or {}).items():
            verified[slug] = bool(info.get("ok"))
        print("read verification report: %d entries (%d failed)"
              % (len(verified), sum(1 for v in verified.values() if not v)))

    local = {}
    if os.path.isdir(args.videos_dir):
        for root, _dirs, files in os.walk(args.videos_dir):
            if os.path.basename(root) == "_failed":
                continue  # quarantined files never count as downloaded
            for fn in files:
                if not fn.lower().endswith((".mp4", ".mkv", ".webm", ".m4v")):
                    continue
                if ".part" in fn:
                    continue
                for v in events:
                    for vv in (v.get("videos") or []):
                        if vv.get("id") and str(vv["id"]) in fn:
                            slug = v.get("slug") or ""
                            # if a report exists, trust it over mere file presence
                            if slug in verified and not verified[slug]:
                                continue
                            rel = os.path.relpath(os.path.join(root, fn), args.videos_dir)
                            local[fn] = (v, rel)

    wb = openpyxl.Workbook()
    wb.properties.title = args.title

    # ------------------------------------------------------------ 明细
    ws = wb.active
    ws.title = "明细"
    heads = [("序号", 6), ("日期", 12), ("标题", 50), ("主题分类", 14), ("演讲人", 32),
             ("是否有录像", 10), ("视频平台", 10), ("视频ID", 12), ("视频链接", 42),
             ("页面链接", 42), ("关联资料", 30), ("简介", 56), ("本地文件", 34), ("是否已下载", 10),
             ("关键词命中", 12)]
    for i, (n, w) in enumerate(heads, 1):
        c = ws.cell(row=1, column=i, value=n)
        c.font = Font(bold=True, color=XL_WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=XL_PRIMARY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30

    rows = sorted(events, key=lambda r: (r.get("date") or "0000-00-00"), reverse=True)
    FIRST = 2
    for idx, r in enumerate(rows):
        row = FIRST + idx
        vlink, vid = video_cell(r)
        plat = (r.get("videos") or [{}])[0].get("platform", "")
        lf = next((rel for fn, (v, rel) in local.items() if v is r), "")
        ws.cell(row=row, column=1, value=idx + 1).alignment = Alignment(horizontal="center")
        dt = parse_dt(r.get("date"), r.get("time"))
        dc = ws.cell(row=row, column=2, value=dt)
        dc.number_format = "YYYY-MM-DD"
        dc.alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=3, value=r.get("title", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row, column=4, value=topic_of(r)).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=5, value="; ".join(r.get("presenters") or [])).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row, column=6, value="是" if r.get("has_video") else "否").alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=7, value=plat).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=8, value=vid).alignment = Alignment(horizontal="center")
        lc = ws.cell(row=row, column=9, value=vlink)
        if vlink.startswith("http"):
            lc.hyperlink = vlink
            lc.style = "Hyperlink"
        pc = ws.cell(row=row, column=10, value=r.get("url", ""))
        pc.hyperlink = r.get("url") or None
        pc.style = "Hyperlink"
        atts = "; ".join(a["name"] for a in (r.get("attachments") or []))
        ws.cell(row=row, column=11, value=atts).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row, column=12, value=r.get("description", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row, column=13, value=lf).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row, column=14, value="是" if lf else "否").alignment = Alignment(horizontal="center")
        # Explicit keyword-match column: lets 汇总 use an EXACT COUNTIF instead of a
        # wildcard, and keeps the 汇总 number identical to the 专题 sheet's filter
        # (which matches on title + description).
        kw_hit = ""
        if kw:
            hay = ((r.get("title") or "") + " " + (r.get("description") or "")).lower()
            if kw.lower() in hay:
                kw_hit = kw
        ws.cell(row=row, column=15, value=kw_hit).alignment = Alignment(horizontal="center")
        for col in range(1, 16):
            ws.cell(row=row, column=col).border = BOX
            if row % 2 == 0 and ws.cell(row=row, column=col).fill.start_color.rgb in (None, "00000000"):
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=XL_NOTE)
    LAST = FIRST + len(rows) - 1
    ws.auto_filter.ref = "A1:O%d" % LAST
    ws.freeze_panes = "C2"

    # ------------------------------------------------------------ 汇总
    sm = wb.create_sheet("汇总", 0)
    sm.sheet_properties.tabColor = XL_PRIMARY
    sm["A1"] = args.title
    sm["A1"].font = Font(bold=True, size=15, color=XL_WHITE)
    sm["A1"].fill = PatternFill("solid", fgColor=XL_DARK)
    sm["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sm.merge_cells("A1:F1")
    sm.row_dimensions[1].height = 34
    for i, w in enumerate([24, 14, 14, 6, 22, 14], 1):
        sm.column_dimensions[get_column_letter(i)].width = w

    kpis = ["活动总场次", "含公开录像场次", "无公开录像场次", "关键词相关场次", "已下载到本地"]
    for i, (name, w) in enumerate(zip(kpis, [24, 14, 14, 6, 22, 14]), 1):
        sm.cell(row=3, column=i, value=name)
    sm["A3"] = "指标"
    sm["B3"] = "数值"
    sm["C3"] = "单位"
    for c in ("A3", "B3", "C3"):
        sm[c].font = Font(bold=True, color=XL_WHITE)
        sm[c].fill = PatternFill("solid", fgColor=XL_PRIMARY)
        sm[c].alignment = Alignment(horizontal="center", vertical="center")
        sm[c].border = BOX
    for i, label in enumerate(kpis):
        r = 4 + i
        sm.cell(row=r, column=1, value=label).font = Font(bold=True, size=10)
        sm.cell(row=r, column=2).font = Font(bold=True, size=12, color=XL_DARK)
        sm.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        sm.cell(row=r, column=3, value="场" if i < 4 else "个").alignment = Alignment(horizontal="center")
        for col in range(1, 4):
            sm.cell(row=r, column=col).border = BOX
            sm.cell(row=r, column=col).fill = PatternFill("solid", fgColor=XL_LIGHT)
    sm["B4"] = "=COUNTA(明细!$C$%d:$C$%d)" % (FIRST, LAST)
    sm["B5"] = '=COUNTIF(明细!$F$%d:$F$%d,"是")' % (FIRST, LAST)
    sm["B6"] = '=COUNTIF(明细!$F$%d:$F$%d,"否")' % (FIRST, LAST)
    if kw:
        # EXACT match on the explicit 关键词命中 column (no wildcard).
        # Same predicate as the 专题 sheet filter, so both numbers agree.
        sm["B7"] = '=COUNTIF(明细!$O$%d:$O$%d,"%s")' % (FIRST, LAST, kw)
    else:
        sm["B7"] = "=0"
    sm["B8"] = '=COUNTIF(明细!$N$%d:$N$%d,"是")' % (FIRST, LAST)

    # year distribution
    years = sorted({(r.get("date") or "")[:4] for r in rows if r.get("date")}, reverse=True)
    sm["A11"] = "年度分布"
    sm["A11"].font = Font(bold=True, size=12, color=XL_WHITE)
    sm["A11"].fill = PatternFill("solid", fgColor=XL_DARK)
    sm.merge_cells("A11:C11")
    for i, n in enumerate(["年份", "场次", "占比"], 1):
        c = sm.cell(row=12, column=i, value=n)
        c.font = Font(bold=True, color=XL_WHITE)
        c.fill = PatternFill("solid", fgColor=XL_PRIMARY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    for i, y in enumerate(years):
        r = 13 + i
        sm.cell(row=r, column=1, value=int(y)).alignment = Alignment(horizontal="center")
        sm.cell(row=r, column=2, value='=COUNTIFS(明细!$B$%d:$B$%d,">="&DATE(%s,1,1),明细!$B$%d:$B$%d,"<="&DATE(%s,12,31))'
                % (FIRST, LAST, y, FIRST, LAST, y)).alignment = Alignment(horizontal="center")
        sm.cell(row=r, column=3, value='=IF(OR(B%d="",B%d=0,$B$4=0),"",B%d/$B$4)' % (r, r, r))
        sm.cell(row=r, column=3).number_format = "0.0%"
        for col in range(1, 4):
            sm.cell(row=r, column=col).border = BOX
    ty = 13 + len(years)
    sm.cell(row=ty, column=1, value="合计")
    sm.cell(row=ty, column=2, value="=SUM(B13:B%d)" % (ty - 1))
    sm.cell(row=ty, column=3, value='=IF(OR(B%d="",B%d=0,$B$4=0),"",B%d/$B$4)' % (ty, ty, ty))
    sm.cell(row=ty, column=3).number_format = "0.0%"
    for col in range(1, 4):
        c = sm.cell(row=ty, column=col)
        c.font = Font(bold=True, color=XL_WHITE)
        c.fill = PatternFill("solid", fgColor=XL_DARK)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX

    # topic distribution
    labels = [t[0] for t in TOPIC_RULES] + ["其他"]
    sm["E11"] = "主题分布"
    sm["E11"].font = Font(bold=True, size=12, color=XL_WHITE)
    sm["E11"].fill = PatternFill("solid", fgColor=XL_DARK)
    sm.merge_cells("E11:F11")
    for i, n in enumerate(["主题分类", "场次"], 5):
        c = sm.cell(row=12, column=i, value=n)
        c.font = Font(bold=True, color=XL_WHITE)
        c.fill = PatternFill("solid", fgColor=XL_PRIMARY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    for i, lab in enumerate(labels):
        r = 13 + i
        sm.cell(row=r, column=5, value=lab)
        sm.cell(row=r, column=6, value='=COUNTIF(明细!$D$%d:$D$%d,E%d)' % (FIRST, LAST, r))
        sm.cell(row=r, column=6).alignment = Alignment(horizontal="center")
        for col in (5, 6):
            sm.cell(row=r, column=col).border = BOX

    try:
        from openpyxl.chart import BarChart, Reference
        ch = BarChart()
        ch.type = "col"
        ch.style = 10
        ch.title = "各年度场次"
        ch.y_axis.title = "场次"
        ch.height, ch.width = 7.5, 15
        ch.add_data(Reference(sm, min_col=2, min_row=12, max_row=12 + len(years)), titles_from_data=True)
        ch.set_categories(Reference(sm, min_col=1, min_row=13, max_row=12 + len(years)))
        sm.add_chart(ch, "E%d" % (13 + len(labels)))
    except Exception:  # noqa: BLE001
        pass

    # ------------------------------------------------------------ 专题
    if kw:
        ss = wb.create_sheet(args.subset_name, 1)
        ss.sheet_properties.tabColor = xl_color("#548235")
        ss["A1"] = "%s 专题" % kw
        ss["A1"].font = Font(bold=True, size=14, color=XL_WHITE)
        ss["A1"].fill = PatternFill("solid", fgColor=xl_color("#548235"))
        ss["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ss.merge_cells("A1:H1")
        ss.row_dimensions[1].height = 30
        cheads = [("序号", 6), ("日期", 12), ("标题", 50), ("是否有录像", 10),
                  ("视频ID", 12), ("视频链接", 42), ("本地文件", 34), ("下载状态", 12)]
        for i, (n, w) in enumerate(cheads, 1):
            c = ss.cell(row=2, column=i, value=n)
            c.font = Font(bold=True, color=XL_WHITE, size=10)
            c.fill = PatternFill("solid", fgColor=xl_color("#548235"))
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BOX
            ss.column_dimensions[get_column_letter(i)].width = w
        sub = [r for r in rows if kw.lower() in ((r.get("title") or "") + " " +
                                                (r.get("description") or "")).lower()]
        for idx, r in enumerate(sub):
            row = 3 + idx
            vlink, vid = video_cell(r)
            lf = next((rel for fn, (v, rel) in local.items() if v is r), "")
            ss.cell(row=row, column=1, value=idx + 1).alignment = Alignment(horizontal="center")
            dt = parse_dt(r.get("date"), r.get("time"))
            dc = ss.cell(row=row, column=2, value=dt)
            dc.number_format = "YYYY-MM-DD"
            dc.alignment = Alignment(horizontal="center")
            ss.cell(row=row, column=3, value=r.get("title", "")).alignment = Alignment(vertical="top", wrap_text=True)
            ss.cell(row=row, column=4, value="是" if r.get("has_video") else "否").alignment = Alignment(horizontal="center")
            ss.cell(row=row, column=5, value=vid).alignment = Alignment(horizontal="center")
            lk = ss.cell(row=row, column=6, value=vlink)
            if vlink.startswith("http"):
                lk.hyperlink = vlink
                lk.style = "Hyperlink"
            ss.cell(row=row, column=7, value=lf).alignment = Alignment(vertical="top", wrap_text=True)
            st = ss.cell(row=row, column=8, value="已下载" if lf else ("无录像" if not r.get("has_video") else "待下载"))
            st.alignment = Alignment(horizontal="center")
            for col in range(1, 9):
                ss.cell(row=row, column=col).border = BOX
            if lf:
                for col in range(1, 9):
                    ss.cell(row=row, column=col).fill = PatternFill("solid", fgColor=XL_OK_BG)
            elif not r.get("has_video"):
                for col in range(1, 9):
                    ss.cell(row=row, column=col).fill = PatternFill("solid", fgColor=XL_WARN_BG)
        ss.auto_filter.ref = "A2:H%d" % (2 + len(sub))
        ss.freeze_panes = "C3"

    # ------------------------------------------------------------ 说明
    sn = wb.create_sheet("下载说明", 2)
    sn.sheet_properties.tabColor = XL_BORDER
    sn.column_dimensions["A"].width = 20
    sn.column_dimensions["B"].width = 100
    sn["A1"] = "清单与下载说明"
    sn["A1"].font = Font(bold=True, size=14, color=XL_WHITE)
    sn["A1"].fill = PatternFill("solid", fgColor=XL_DARK)
    sn["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sn.merge_cells("A1:B1")
    sn.row_dimensions[1].height = 30
    for i, n in enumerate(["项目", "内容"], 1):
        c = sn.cell(row=3, column=i, value=n)
        c.font = Font(bold=True, color=XL_WHITE)
        c.fill = PatternFill("solid", fgColor=XL_PRIMARY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    doc = [("清单来源", data.get("listing", "")),
           ("活动总数", str(len(rows))),
           ("含录像", str(sum(1 for r in rows if r.get("has_video")))),
           ("抓取时间", datetime.now().strftime("%Y-%m-%d %H:%M"))]
    if args.notes and os.path.exists(args.notes):
        with open(args.notes, encoding="utf-8") as f:
            extra = [ln.strip() for ln in f if ln.strip()]
        doc += [(ln.split("：")[0] if "：" in ln else "备注", ln) for ln in extra]
    for i, (k, v) in enumerate(doc):
        r = 4 + i
        sn.cell(row=r, column=1, value=k).font = Font(bold=True, size=10)
        sn.cell(row=r, column=2, value=v).alignment = Alignment(vertical="top", wrap_text=True)
        for col in (1, 2):
            sn.cell(row=r, column=col).border = BOX
            sn.cell(row=r, column=col).fill = PatternFill("solid", fgColor=XL_LIGHT if i % 2 else XL_NOTE)
    sn.freeze_panes = "A4"

    out = os.path.abspath(args.out)
    wb.save(out)
    print("saved:", out)
    print("sheets:", wb.sheetnames, "| detail rows:", len(rows))


if __name__ == "__main__":
    main()
